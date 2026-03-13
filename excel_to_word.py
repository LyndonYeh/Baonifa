import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from collections import defaultdict
import os
import win32com.client
import re

# 設定日誌層級 ('DEBUG', 'INFO', 'WARNING', 'ERROR')
LOG_LEVEL = "WARNING"

def log(level, message):
    """根據 LOG_LEVEL 顯示對應層級的日誌"""
    levels = ["DEBUG", "INFO", "WARNING", "ERROR"]
    if levels.index(level) >= levels.index(LOG_LEVEL):
        print(f"[{level}] {message}")

def wrap_text_in_cell(cell, text, max_width):
    """將文字自動換行（支援 str 或 list），並逐行插入段落"""
    lines = []

    if isinstance(text, list):
        lines = text
    else:
        words = text.split()
        current_line = ""
        for word in words:
            if len(current_line) + len(word) + 1 > max_width:
                lines.append(current_line)
                current_line = word
            else:
                if current_line:
                    current_line += " "
                current_line += word
        if current_line:
            lines.append(current_line)

    if not lines:
        cell.text = ""
        return

    first_p = cell.paragraphs[0]
    first_p.text = lines[0]
    first_p.paragraph_format.space_after = Pt(0)

    for line in lines[1:]:
        p = cell.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)


def format_path(path):
    """將修改清單中的 `/` 轉換為 `_`，用於匹配 diff_patch 資料夾內的 patch 檔案"""
    return path.replace("/", "_")

def clean_path(path):
    """移除非預期的控制字符，如 `\x07`"""
    return re.sub(r'[^\x20-\x7E]', '', path).strip()

def extract_file_paths_from_table(doc):
    """從每張表格的修改清單欄位提取所有檔案路徑"""
    tables = doc.Tables
    attachments_map = {}

    for table_index in range(1, len(tables) + 1):
        try:
            table = tables(table_index)
            mod_list_cell = table.Cell(3, 2)
            raw_text = mod_list_cell.Range.Text.strip()

            lines = re.split(r'[\r\n\x0b]+', raw_text)
            file_paths = []
            for line in lines:
                file_paths.extend([
                    clean_path(p) for p in line.strip().split()
                    if p and p.strip() not in [".", "", "diff_patch/.patch"]
                ])

            patch_files = [
                f"diff_patch/{format_path(path)}.patch"
                for path in file_paths
                if path.strip() and not path.endswith(".patch")
            ]

            attachments_map[table_index] = patch_files
        except Exception as e:
            log("ERROR", f"Extracting file paths from table {table_index}: {e}")
            continue

    return attachments_map

def get_filename_from_path(path):
    """取得檔案名稱，從最後一個 `_` 之後的部分"""
    filename = os.path.basename(path).replace(".patch", "")
    return re.split(r'[_]', filename)[-1] if re.search(r'[_]', filename) else filename

def insert_attachments_to_tables(doc_path, base_dir=None):
    """使用 win32com 將檔案內嵌到 Word 表格內的正確格子"""
    if base_dir is None:
        base_dir = os.getcwd()

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(os.path.abspath(doc_path))

    attachments_map = extract_file_paths_from_table(doc)
    tables = doc.Tables

    for table_index, patch_files in attachments_map.items():
        if table_index <= len(tables):
            table = tables(table_index)
            cell = table.Cell(4, 2)
            cell.Range.Text = ""

            for patch_file in patch_files:
                filename = get_filename_from_path(patch_file)
                abs_patch = os.path.normpath(os.path.join(base_dir, patch_file))
                if os.path.exists(abs_patch):
                    log("INFO", f"Inserting: {abs_patch} into table {table_index}")
                    try:
                        shape = cell.Range.InlineShapes.AddOLEObject(
                            ClassType="Package",
                            FileName=abs_patch,
                            LinkToFile=False,
                            DisplayAsIcon=True,
                        )
                        shape.OLEFormat.IconLabel = filename
                        shape.Width = 50
                        shape.Height = 50
                    except Exception as e:
                        log("ERROR", f"OLE 插入失敗：{abs_patch} - {e}")
                else:
                    log("WARNING", f"File not found: {abs_patch}")

    doc.Save()
    doc.Close()
    word.Quit()

def excel_to_word(excel_file, word_file, base_dir=None):
    df = pd.read_excel(excel_file)
    df = df.fillna("").astype(str)
    filtered_df = df[df["過版內容說明"].str.startswith(("fix:", "feat", "refactor"), na=False)]

    grouped_data = defaultdict(lambda: {"模組代號": set(), "模組名稱": set(), "修改清單": []})

    for _, row in filtered_df.iterrows():
        d_value = row["過版內容說明"]
        original_path = f"{row['程式路徑']}/{row['檔名']}" if row["程式路徑"] else row["檔名"]
        grouped_data[d_value]["模組代號"].add(row["模組代號"])
        grouped_data[d_value]["模組名稱"].add(row["模組名稱"])
        if original_path not in grouped_data[d_value]["修改清單"]:
            grouped_data[d_value]["修改清單"].append(original_path)

    doc = Document()
    doc.add_heading("程式修改說明", level=1)

    for d_value, content in grouped_data.items():
        table = doc.add_table(rows=4, cols=4, style="Table Grid")
        table.autofit = False

        table.cell(0, 0).text = "模組代號"
        table.cell(0, 1).text = "、".join(content["模組代號"])
        table.cell(0, 2).text = "模組名稱"
        table.cell(0, 3).text = "、".join(content["模組名稱"])

        table.cell(1, 0).text = "修改目的"
        table.cell(1, 1).merge(table.cell(1, 3))
        wrap_text_in_cell(table.cell(1, 1), d_value, max_width=70)

        table.cell(2, 0).text = "修改清單"
        table.cell(2, 1).merge(table.cell(2, 3))
        wrap_text_in_cell(table.cell(2, 1), content["修改清單"], max_width=70)

        table.cell(3, 0).text = "附件"
        table.cell(3, 1).merge(table.cell(3, 3))

        doc.add_paragraph("")

    doc.save(word_file)
    insert_attachments_to_tables(word_file, base_dir=base_dir)


# 使用
if __name__ == "__main__":
    excel_to_word("程式修改清單.xlsx", "程式修改說明.docx")
