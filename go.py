import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import threading
import queue
import subprocess
import os
import io
import re
import shutil
import zipfile
import traceback
import sys

_NO_WINDOW = subprocess.CREATE_NO_WINDOW if hasattr(subprocess, "CREATE_NO_WINDOW") else 0


def resource_path(filename):
    """取得資源檔路徑，相容 PyInstaller 打包後的環境"""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, filename)


def _import_scripts():
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from patch_to_excel import parse_patch_to_excel
    from excel_to_word import excel_to_word
    return parse_patch_to_excel, excel_to_word


def get_version(repo_path, target):
    """從 target commit 的 merge message 取得版本名稱（如 20260313）"""
    result = subprocess.run(
        ["git", "log", "-1", "--format=%s", target],
        capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=repo_path,
        creationflags=_NO_WINDOW,
    )
    msg = result.stdout.strip()
    # 匹配 "Merge branch 'temp/20260313' into ..." 格式
    match = re.search(r"branch\s+'[^']*?/([^/']+)'", msg)
    if match:
        return match.group(1)
    # fallback: 找 8 位數日期
    match = re.search(r'\b(\d{8})\b', msg)
    if match:
        return match.group(1)
    # 最後 fallback: short commit hash
    return target[:8]


def fetch_commits(repo_path):
    """抓取 source（遠端 develop）和 target（本地 develop）的最新 commit hash"""
    try:
        r = subprocess.run(["git", "rev-parse", "develop"],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=repo_path,
                           creationflags=_NO_WINDOW,
                           )
        target = r.stdout.strip() if r.returncode == 0 else ""
    except Exception:
        target = ""
    try:
        r = subprocess.run(["git", "ls-remote", "origin", "refs/heads/develop"],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=repo_path,
                           creationflags=_NO_WINDOW,
                           )
        source = r.stdout.split()[0] if r.returncode == 0 and r.stdout.strip() else ""
    except Exception:
        source = ""
    return source, target


def worker(repo_path, source, target, output_dir, q):
    def log(msg):
        q.put(("log", msg))

    def progress(val):
        q.put(("progress", val))

    try:
        # ── 取得版本名稱 ──────────────────────────────────────────
        version = get_version(repo_path, target)
        date_str = f"{version[:4]}/{version[4:6]}/{version[6:8]}" if len(version) == 8 and version.isdigit() else version
        log(f"上版日期：{date_str}")

        diff_patch_dir  = os.path.join(output_dir, "diff_patch")
        diff_source_dir = os.path.join(output_dir, "diff_source")
        diff_patch_file = os.path.join(output_dir, "diff.patch")
        diff_delete_file = os.path.join(output_dir, "diff_delete.txt")
        excel_file      = os.path.join(output_dir, "程式修改清單.xlsx")
        word_file       = os.path.join(output_dir, "程式修改說明.docx")
        source_zip_name = f"{version}-差異檔.zip"
        source_zip_path = os.path.join(output_dir, source_zip_name)
        final_zip_name  = f"{version}上版清單.zip"
        final_zip_path  = os.path.join(output_dir, final_zip_name)

        os.makedirs(diff_patch_dir, exist_ok=True)

        # ── Step 1：產生 diff.patch ──────────────────────────────
        log(f"[1/8] 產生所有 commit 的完整差異")
        with open(diff_patch_file, "w", encoding="utf-8") as f:
            result = subprocess.run(
                ["git", "log", "-p", f"{source}..{target}"],
                stdout=f, stderr=subprocess.PIPE, cwd=repo_path,
                creationflags=_NO_WINDOW,
            )
        if result.returncode != 0:
            raise RuntimeError(f"git log 失敗：{result.stderr.decode('utf-8', errors='replace')}")
        progress(10)

        # ── Step 2：取得變更清單，逐一產生 diff_patch/*.patch ─────
        log("[2/8] 每個 diff_patch 各自產出")
        result = subprocess.run(
            ["git", "diff", "--name-only", "--diff-filter=ACMRT", source, target],
            capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=repo_path,
            creationflags=_NO_WINDOW,
        )
        if result.returncode != 0:
            raise RuntimeError(f"git diff --name-only 失敗：{result.stderr}")

        changed_files = [f for f in result.stdout.splitlines() if f.strip()]
        log(f"    共 {len(changed_files)} 個變更檔案")

        for file_path in changed_files:
            patch_name = file_path.replace("/", "_") + ".patch"
            patch_out = os.path.join(diff_patch_dir, patch_name)
            diff_result = subprocess.run(
                ["git", "diff", source, target, "--", file_path],
                capture_output=True, cwd=repo_path,
                creationflags=_NO_WINDOW,
            )
            with open(patch_out, "wb") as f:
                f.write(diff_result.stdout)
        progress(25)

        # ── Step 3：git archive → diff_source/ → {version}-差異檔.zip ─
        if changed_files:
            log(f"[3/8] 打包 source 資料夾")
            archive_result = subprocess.run(
                ["git", "archive", "--format=zip", target] + changed_files,
                capture_output=True, cwd=repo_path,
                creationflags=_NO_WINDOW,
            )
            if archive_result.returncode != 0:
                raise RuntimeError(f"git archive 失敗：{archive_result.stderr.decode('utf-8', errors='replace')}")
            os.makedirs(diff_source_dir, exist_ok=True)
            with zipfile.ZipFile(io.BytesIO(archive_result.stdout)) as zf:
                zf.extractall(diff_source_dir)
            # 壓縮成 {version}-差異檔.zip
            with zipfile.ZipFile(source_zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
                for root, _, files in os.walk(diff_source_dir):
                    for fname in files:
                        fp = os.path.join(root, fname)
                        zf.write(fp, os.path.relpath(fp, diff_source_dir))
        else:
            log("[3/8] 無變更檔案，略過差異檔步驟")
        progress(45)

        # ── Step 4：產生 diff_delete.txt ─────────────────────────
        log("[4/8] 產出待刪除列表")
        del_result = subprocess.run(
            ["git", "diff", "--name-only", "--diff-filter=D", source, target],
            capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=repo_path,
            creationflags=_NO_WINDOW,
        )
        if del_result.returncode != 0:
            raise RuntimeError(f"git diff --diff-filter=D 失敗：{del_result.stderr}")
        with open(diff_delete_file, "w", encoding="utf-8") as f:
            f.write(del_result.stdout)
        has_deletes = bool(del_result.stdout.strip())
        log(f"    {'有' if has_deletes else '無'}刪除檔案")
        progress(55)

        # ── Step 5：patch → Excel ─────────────────────────────────
        log("[5/8] 產出 excel")
        parse_patch_to_excel, excel_to_word = _import_scripts()
        parse_patch_to_excel(diff_patch_file, excel_file)

        # 填入過版日期
        from openpyxl import load_workbook
        wb = load_workbook(excel_file)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, max_col=1):
            for cell in row:
                cell.value = version
        wb.save(excel_file)
        progress(70)

        # ── Step 6：Excel → Word（含嵌入 patch 附件）─────────────
        log("[6/8] 將 excel 轉 word（含嵌入檔案附件）")

        class _StdoutToLog:
            def write(self, text):
                if text.strip():
                    q.put(("log", f"    {text.strip()}"))
            def flush(self):
                pass

        _old_stdout = sys.stdout
        sys.stdout = _StdoutToLog()
        try:
            excel_to_word(excel_file, word_file, base_dir=output_dir)
        finally:
            sys.stdout = _old_stdout
        progress(82)

        # ── Step 7：打包最終上版清單 ──────────────────────────────
        log("[7/8] 將 source & excel & word 打包進上版清單")
        with zipfile.ZipFile(final_zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            if os.path.exists(source_zip_path):
                zf.write(source_zip_path, source_zip_name)
            if os.path.exists(excel_file):
                zf.write(excel_file, "程式修改清單.xlsx")
            if os.path.exists(word_file):
                zf.write(word_file, "程式修改說明.docx")
            if has_deletes and os.path.exists(diff_delete_file):
                zf.write(diff_delete_file, "待刪除名單.txt")
        progress(93)

        # ── Step 8：清理中間產物 ──────────────────────────────────
        log("[8/8] 清理所有暫存檔")
        for item in [diff_patch_file, diff_patch_dir, diff_source_dir,
                     diff_delete_file, excel_file, word_file, source_zip_path]:
            try:
                if os.path.isdir(item):
                    shutil.rmtree(item)
                elif os.path.exists(item):
                    os.remove(item)
            except Exception as e:
                log(f"    清理略過：{os.path.basename(item)} ({e})")
        progress(100)

        q.put(("done", final_zip_path))

    except Exception:
        q.put(("error", traceback.format_exc()))


def create_gui():
    root = tk.Tk()
    root.title("Baonifa")
    root.resizable(False, False)

    icon_path = resource_path("robot.ico")
    if os.path.exists(icon_path):
        root.iconbitmap(icon_path)

    pad = {"padx": 8, "pady": 4}

    input_frame = ttk.LabelFrame(root, text="設定", padding=8)
    input_frame.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 4))

    repo_var   = tk.StringVar(value="")
    output_var = tk.StringVar(value="")
    source_var = tk.StringVar()
    target_var = tk.StringVar()

    def browse_dir(var):
        path = filedialog.askdirectory()
        if path:
            var.set(path)

    def refresh_commits():
        repo = repo_var.get().strip()
        if not repo:
            return
        src, tgt = fetch_commits(repo)
        source_var.set(src or "（無法取得）")
        target_var.set(tgt or "（無法取得）")

    rows = [
        ("Git Repo 路徑：", repo_var,   True),
        ("輸出資料夾：",    output_var,  True),
        ("Source Commit：", source_var, False),
        ("Target Commit：", target_var, False),
    ]
    for i, (label, var, has_browse) in enumerate(rows):
        ttk.Label(input_frame, text=label, width=14, anchor="e").grid(row=i, column=0, **pad)
        ttk.Entry(input_frame, textvariable=var, width=45).grid(row=i, column=1, **pad)
        if has_browse:
            ttk.Button(input_frame, text="瀏覽", width=5,
                       command=lambda v=var: browse_dir(v)).grid(row=i, column=2, **pad)

    # 啟動時自動抓取 commit
    root.after(100, refresh_commits)

    generate_btn = ttk.Button(root, text="產生上版清單", width=20)
    generate_btn.grid(row=1, column=0, pady=8)

    log_frame = ttk.LabelFrame(root, text="執行記錄", padding=4)
    log_frame.grid(row=2, column=0, sticky="ew", padx=12, pady=4)
    log_text = scrolledtext.ScrolledText(log_frame, width=62, height=12, state="disabled",
                                         font=("Consolas", 9))
    log_text.pack()

    progress_var = tk.IntVar(value=0)
    ttk.Progressbar(root, variable=progress_var, maximum=100, length=480).grid(
        row=3, column=0, padx=12, pady=(4, 12))

    def append_log(msg):
        log_text.config(state="normal")
        log_text.insert(tk.END, msg + "\n")
        log_text.see(tk.END)
        log_text.config(state="disabled")

    q = queue.Queue()

    def poll_queue():
        while not q.empty():
            msg_type, payload = q.get_nowait()
            if msg_type == "log":
                append_log(payload)
            elif msg_type == "progress":
                progress_var.set(payload)
            elif msg_type == "done":
                append_log(f"\n完成！上版清單：{payload}")
                generate_btn.config(state="normal")
                messagebox.showinfo("完成", f"上版清單已產生：\n{payload}")
            elif msg_type == "error":
                append_log(f"\n[錯誤]\n{payload}")
                generate_btn.config(state="normal")
                messagebox.showerror("執行失敗", payload.splitlines()[-1])
        root.after(100, poll_queue)

    def on_generate():
        repo_path  = repo_var.get().strip()
        source     = source_var.get().strip()
        target     = target_var.get().strip()
        output_dir = output_var.get().strip()

        errors = []
        if not repo_path:
            errors.append("請填入 Git Repo 路徑")
        elif not os.path.isdir(repo_path):
            errors.append(f"Repo 路徑不存在：{repo_path}")
        if not source or source == "（無法取得）":
            errors.append("無法取得 Source Commit，請確認遠端連線是否正常")
        if not target or target == "（無法取得）":
            errors.append("無法取得 Target Commit，請確認 develop 分支是否存在")
        if not output_dir:
            errors.append("請填入輸出資料夾")
        elif not os.path.isdir(output_dir):
            errors.append(f"輸出資料夾不存在：{output_dir}")

        if errors:
            messagebox.showerror("驗證失敗", "\n".join(errors))
            return

        log_text.config(state="normal")
        log_text.delete("1.0", tk.END)
        log_text.config(state="disabled")
        progress_var.set(0)

        generate_btn.config(state="disabled")
        threading.Thread(
            target=worker,
            args=(repo_path, source, target, output_dir, q),
            daemon=True
        ).start()

    generate_btn.config(command=on_generate)
    root.after(100, poll_queue)
    root.mainloop()


if __name__ == "__main__":
    create_gui()
